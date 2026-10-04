#!/usr/bin/env python3
from docx import Document
doc = Document()
doc.add_heading('Example Assignment', level=0)
doc.add_paragraph('This is an example assignment for testing the local-lab-orchestrator.')
doc.add_paragraph('Task: Write a script that prints "Hello, World!".')
doc.add_paragraph('Required sections:')
doc.add_paragraph('1. Task description')
doc.add_paragraph('2. Variant (if any)')
doc.add_paragraph('3. Solution')
doc.add_paragraph('4. Code')
doc.add_paragraph('5. Conclusion')
doc.save('example_assignment.docx')
print("Created example_assignment.docx")